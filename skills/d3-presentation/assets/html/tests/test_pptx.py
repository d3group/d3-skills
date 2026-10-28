import os, re, subprocess, sys, tempfile, unittest, zipfile
HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)
from d3deck import deck as B

try:
    import playwright  # noqa: F401
    HAVE_PW = True
except ImportError:
    HAVE_PW = False
try:
    import pptx  # noqa: F401
    from d3deck import pptx_export as X
    HAVE_PPTX = True
except ImportError:
    HAVE_PPTX = False


def build(look='klar'):
    B.meta(title='Export test', subtitle='Sub', presenter='A. Author', sections=['Part one', 'Part two'], look=look)
    B.titleslide()
    B.section(1, question='What moves?')
    B.onecol(B.K('Kick', 'A list with one build'), '<ul><li>always there</li><li data-s="1">second step</li></ul>',
             steps=2, notes='Say <b>this</b>.')
    B.onecol('Labels between formulas', '<div><div class="lab">Objective</div>$\\min c^\\top x$<div class="lab">Limits</div>'
             'text with $x \\ge 0$ inline</div>')
    B.section(2)
    B.onecol('Quote', '<div class="stmt"><span>Words<cite style="display:block">Someone (1979)</cite></span></div>')
    B.thankyou()
    out = os.path.join(tempfile.mkdtemp(), 'deck.html')
    B.build('deck', out=out, embed_fonts=False)
    return out


def export(deck, *args):
    out = os.path.join(tempfile.mkdtemp(), 'deck.pptx')
    r = subprocess.run([sys.executable, os.path.join(HERE, 'd3deck', 'pptx_export.py'), deck, '--out', out, *args],
                       capture_output=True, text=True)
    return r, out


def uni_master():
    path = os.environ.get('D3_UNI_MASTER')
    return path if path and os.path.exists(path) else None


def slide_xml(path):
    with zipfile.ZipFile(path) as z:
        names = sorted((n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)),
                       key=lambda n: int(re.search(r'\d+', n.rsplit('/', 1)[1]).group()))
        return [z.read(n).decode() for n in names], z.read('ppt/presentation.xml').decode()


@unittest.skipUnless(HAVE_PPTX, 'python-pptx not installed')
class Runs(unittest.TestCase):

    def test_whitespace_collapses_like_html(self):
        out = X.Writer.clean_runs([{'t': '  two   words '}, {'t': ' more\n text  '}])
        self.assertEqual(''.join(r['t'] for r in out), 'two words more text')

    def test_pre_keeps_lines_as_breaks(self):
        out = X.Writer.clean_runs([{'t': 'a = 1\n  b = 2', 'pre': True}])
        self.assertEqual([r.get('t', 'BR') for r in out], ['a = 1', 'BR', '  b = 2'])

    def test_marker_tab_is_not_trimmed(self):
        out = X.Writer.clean_runs([{'t': 'A', 'pre': True}, {'t': '\t', 'pre': True}, {'t': ' yes'}])
        self.assertEqual(''.join(r['t'] for r in out), 'A\tyes')


# The export needs the university master, which is not part of the (public) repository: set D3_UNI_MASTER.
@unittest.skipUnless(HAVE_PW and HAVE_PPTX and uni_master(), 'needs playwright, python-pptx and D3_UNI_MASTER')
class Export(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.r, cls.out = export(build())
        cls.slides, cls.pres = slide_xml(cls.out) if cls.r.returncode == 0 else ([], '')

    def test_runs_and_writes_every_slide(self):
        self.assertEqual(self.r.returncode, 0, self.r.stdout + self.r.stderr)
        self.assertEqual(len(self.slides), 7)

    def test_title_goes_into_the_title_placeholder_with_the_kicker(self):
        from pptx import Presentation
        s = Presentation(self.out).slides[2]
        self.assertEqual(s.shapes.title.text_frame.text, 'Kick A list with one build')

    def test_notes_are_plain_text(self):
        from pptx import Presentation
        self.assertEqual(Presentation(self.out).slides[2].notes_slide.notes_text_frame.text, 'Say this.')

    def test_step_build_becomes_a_well_formed_click_animation(self):
        xml = self.slides[2]
        self.assertIn('<p:timing>', xml)
        self.assertIn('<p:tgtEl><p:spTgt', xml)                  # PowerPoint repairs files without the tgtEl wrapper
        self.assertIn('<p:pRg st="1" end="1"/>', xml)            # the second bullet, not the whole list
        ids = [int(i) for i in re.findall(r'<p:cTn id="(\d+)"', xml)]
        self.assertEqual(ids, list(range(1, len(ids) + 1)))      # in document order

    def test_labels_and_text_between_blocks_stay_apart(self):
        xml = self.slides[3]
        texts = re.findall(r'<a:t>([^<]*)</a:t>', xml)
        self.assertIn('Objective', texts)
        self.assertIn('Limits', texts)
        self.assertNotIn('Objective', ''.join(t for t in texts if 'inline' in t))
        self.assertIn('<p:pic>', xml)                            # the formula-only line is a picture

    def test_block_inside_inline_is_kept(self):
        self.assertIn('Someone (1979)', self.slides[5])

    def test_klar_keeps_the_d3_logo_on_the_master_frame(self):
        from pptx import Presentation
        s = Presentation(self.out).slides[2]
        self.assertEqual(s.slide_layout.name, 'Inhaltsfolie 1')
        logos = [sh for sh in s.shapes if sh.shape_type == 13 and sh.left > 9000000 and sh.top < 1000000]   # 13: picture
        self.assertEqual(len(logos), 1)

    def test_sections_follow_the_deck(self):
        names = re.findall(r'<p14:section name="([^"]+)"', self.pres)
        self.assertEqual(names, ['Start', 'Part one', 'Part two'])


@unittest.skipUnless(HAVE_PW and HAVE_PPTX and uni_master(), 'needs playwright, python-pptx and D3_UNI_MASTER')
class UniExport(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.r, cls.out = export(build('uni'), '--master', uni_master())

    def test_slides_use_the_university_layouts(self):
        self.assertEqual(self.r.returncode, 0, self.r.stdout + self.r.stderr)
        from pptx import Presentation
        layouts = [s.slide_layout.name for s in Presentation(self.out).slides]
        self.assertEqual(layouts[0][:9], 'Deckblatt')
        self.assertEqual(layouts[1], 'Trennseite 2')
        self.assertEqual(layouts[-1], 'Abschlussseite')

    def test_sample_slides_and_their_sections_are_gone(self):
        from pptx import Presentation
        self.assertEqual(len(Presentation(self.out).slides), 7)
        _, pres = slide_xml(self.out)
        self.assertNotIn('Anleitung', pres)

    def test_missing_master_names_the_ways_to_provide_it(self):
        env = dict(os.environ, D3_UNI_MASTER='')
        deck = build('uni')
        r = subprocess.run([sys.executable, os.path.join(HERE, 'd3deck', 'pptx_export.py'), deck, '--out',
                            os.path.join(tempfile.mkdtemp(), 'x.pptx')], capture_output=True, text=True, env=env,
                           cwd=tempfile.gettempdir())
        self.assertNotEqual(r.returncode, 0)
        self.assertIn('D3_UNI_MASTER', r.stdout + r.stderr)


if __name__ == '__main__':
    unittest.main()
