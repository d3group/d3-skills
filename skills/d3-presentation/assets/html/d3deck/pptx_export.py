#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Export a built deck to an editable PowerPoint file.

    uv run --with playwright --with python-pptx d3deck/pptx_export.py dist/my_talk.html            # -> dist/my_talk.pptx
    uv run --with playwright --with python-pptx d3deck/pptx_export.py dist/my_talk.html --master ~/PPT-Master-Universitaet-Wuerzburg.potx

The browser lays the deck out; this script measures every slide in its final step
and rebuilds it from native PowerPoint objects at the same positions: titles go
into the layout's title placeholder, paragraphs and lists into text boxes (bullets,
bold, colour and inline math as text), tables into tables, boxes, chevrons and bars
into shapes, and step builds into appear-on-click animations. Only what has no
native counterpart becomes a picture with a transparent background: figures, charts,
form controls, and lines that hold nothing but a formula.

Every look is built on the university frame, so the slides go onto the layouts of
the university master (PPT-Master-Universitaet-Wuerzburg.potx: Deckblatt, Agenda,
Inhaltsfolie, Blanko, Trennseite, Abschlussseite) and fill their placeholders; logo,
rule, footer, page number and seal come from the master, the text takes its theme
font, and the D3 logo of klar and beamer stays on the slide as a picture. The .potx
is not part of the skill; see find_master().

Slides added in PowerPoint look like the rest, speaker notes go into the
notes pages, the deck's sections become PowerPoint sections, and data-jump links
become click actions. The deck is measured in its own font and shown in the master's
(Segoe UI), so a line can wrap at a different word.
"""
from __future__ import annotations

import argparse
import base64
import io
import os
import re
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
PX = 9525                      # EMU per CSS px: the 1280x720 canvas is 13.333 x 7.5 in
MONO = 'Consolas'
SCALE = 3                      # device pixels per CSS px for pictures
SETTLE_MS = 480

# What the layouts of the university master draw themselves, per slide kind: (left out entirely, box left out
# but text kept). Every look shares this frame; klar and beamer add the D3 logo, which stays on the slide.
CHROME = {
    'content': (['.uw-logo', '.uw-rule', '.uw-num', '.uw-foot'], ['.take']),
    'title': (['.uw-logo', '.uw-seal:not(.uw-seal-end)', '.uw-band', '.uw-box:not(.uw-contact)', '.uw-photo',
               '.uw-contact > b', '.uw-web', '.uw-rule'], ['.uw-cover', '.uw-dark', '.uw-contact']),
}
ROLES = [['.uw-title', 'title'], ['.uw-ttitle > span', 'title'], ['.uw-tsub', 'sub'], ['.uw-tname', 'name'],
         ['.uw-tdate', 'date'], ['.uw-st', 'title'], ['.uw-sn', 'num'], ['.uw-agn', 'agn'], ['.uw-agt', 'agt'],
         ['.uw-thx', 'thx'], ['.uw-contact > div', 'contact'], ['.take', 'take'], ['.d3-co', 'cobrand']]
# placeholder idx per role on each layout of PPT-Master-Universitaet-Wuerzburg.potx
UNI_PH = {
    'Deckblatt': {'title': [0], 'sub': [1], 'name': [12], 'date': [13]},
    'Trennseite 2': {'title': [0], 'num': [13]},
    'Agenda 1': {'agn': [13, 15, 19, 20, 22, 26], 'agt': [18, 16, 14, 21, 23, 27]},
    'Inhaltsfolie 1': {'title': [0], 'body': [1]},
    'Inhaltsfolie 2': {'title': [0], 'take': [18]},
    'Blanko 1': {'title': [0]},
    'Quellenangaben': {'body': [19]},
    'Abschlussseite': {'thx': [21], 'contact': [22, 23, 24]},
}
UNI_COVER = {'navy': 'Deckblatt 1', 'photo': 'Deckblatt 2', 'band': 'Deckblatt 4'}
# single-line text the HTML cuts with an ellipsis (footer, source line) is cut the same way: PowerPoint has no ellipsis
CUT_JS = '''() => { for (const el of document.querySelectorAll('.sl.on *')) {
  if (el.children.length || el.scrollWidth <= el.clientWidth + 1) continue;
  const cs = getComputedStyle(el); if (cs.textOverflow !== 'ellipsis' || cs.whiteSpace !== 'nowrap') continue;
  const full = el.textContent.trim(); let lo = 0, hi = full.length;
  while (lo < hi) { const m = (lo + hi + 1) >> 1; el.textContent = full.slice(0, m).trimEnd() + '…';
    if (el.scrollWidth <= el.clientWidth + 1) lo = m; else hi = m - 1; }
  el.textContent = full.slice(0, lo).trimEnd() + '…'; } }'''
INFO_JS = '''() => { const s = document.querySelector('.sl.on'), q = c => s.querySelector(c);
  const cov = q('.uw-cover'); const f = q('.uw-foot');
  // a footer the HTML cuts with an ellipsis is cut the same way, so it does not wrap onto a second line
  const foot = !f ? '' : f.scrollWidth <= f.clientWidth + 1 ? f.textContent.trim()
             : f.textContent.trim().slice(0, Math.floor(f.textContent.trim().length * f.clientWidth / f.scrollWidth) - 2).trimEnd() + '…';
  // the layout's info box only stands in for a takeaway in the frame's slot, not for one placed inside the body
  const tk = q('.take'), tr = tk && tk.getBoundingClientRect();
  const slot = !!tr && Math.abs(tr.left - 74) < 8 && Math.abs(tr.bottom - 662) < 8 && Math.abs(tr.width - 902) < 8;
  return {cover: cov ? (cov.className.match(/uw-c-(\\w+)/) || [])[1] : null, take: slot, refs: !!q('.refs'),
          num: !!q('.uw-num'), foot: foot, footw: f ? f.getBoundingClientRect().width : 0,
          agenda: !!q('.uw-agl'), dark: !!q('.uw-dark'),
          title: (q('.uw-title') || {textContent: ''}).textContent.trim()}; }'''
def E(px):
    return Emu(int(round(px * PX)))


def rgb(c):
    return RGBColor(*(int(round(v)) for v in c[:3]))


def hexc(c):
    return '%02X%02X%02X' % tuple(int(round(v)) for v in c[:3])


def strip_html(s):
    s = re.sub(r'<br\s*/?>|</p>|</li>', '\n', s or '')
    s = re.sub(r'<[^>]+>', '', s)
    for a, b in (('&nbsp;', ' '), ('&amp;', '&'), ('&lt;', '<'), ('&gt;', '>'), ('&mdash;', '—'), ('&rarr;', '→')):
        s = s.replace(a, b)
    return re.sub(r'\n{3,}', '\n\n', s).strip()


# ── the browser side ─────────────────────────────────────────────────────

def measure(deck):
    """-> (manifest, [per slide: {'items', 'pics', 'info'}])."""
    from playwright.sync_api import sync_playwright
    with open(os.path.join(HERE, 'extract.js'), encoding='utf-8') as fh:
        js = fh.read()
    slides = []
    with sync_playwright() as p:
        try:
            browser = p.chromium.launch()
        except Exception as exc:
            sys.exit(f'Playwright could not launch Chromium ({exc.__class__.__name__}). '
                     'Run: uv run --with playwright playwright install chromium')
        page = browser.new_page(viewport={'width': 1280, 'height': 720}, device_scale_factor=SCALE)
        page.goto('file://' + os.path.abspath(deck) + '?shot')
        page.wait_for_timeout(1500)
        page.wait_for_function('window.D3_MATH_READY !== false', timeout=30000)
        # finish transitions and keyframe animations at once: kit objects grow in from an invisible start state
        page.add_style_tag(content='*,*::before,*::after{transition-duration:0s!important;transition-delay:0s!important;'
                                   'animation-duration:0s!important;animation-delay:0s!important;animation-iteration-count:1!important}'
                                   'html.xshot,html.xshot body,html.xshot #stage,html.xshot .sl{background:transparent!important;box-shadow:none!important}'
                                   'html.xshot body *{visibility:hidden!important}'
                                   'html.xshot .xshot-on,html.xshot .xshot-on *{visibility:visible!important}')
        manifest = page.evaluate('META')

        def grab(i, kind):
            ch, bo = CHROME[kind]
            page.evaluate(f'show({i}, steps({i}) - 1)')
            page.wait_for_timeout(60)
            page.evaluate(CUT_JS)
            info = page.evaluate(INFO_JS)
            if info['agenda'] and not info['dark'] and kind == 'content':
                ch = ch + ['.uw-title']                 # 'Agenda' is printed on the layout
            roles = ROLES
            if not info['take']:                        # draw the box and arrow of a takeaway outside the slot
                bo, roles = [b for b in bo if b != '.take'], [r for r in ROLES if r[1] != 'take']
            items = page.evaluate(js, {'mode': 'content', 'chrome': ch, 'boxOnly': bo, 'roles': roles})
            pics = {}
            for it in items:
                if it.get('xid'):
                    # only the element itself: page, slide and the boxes behind it are rebuilt as shapes
                    sel = f'.sl.on [data-xid="{it["xid"]}"]'
                    page.evaluate('''sel => { document.documentElement.classList.add('xshot');
                        document.querySelector(sel).classList.add('xshot-on'); }''', sel)
                    pics[it['xid']] = page.locator(sel).screenshot(omit_background=True)
                    page.evaluate('''() => { document.documentElement.classList.remove('xshot');
                        document.querySelectorAll('.xshot-on').forEach(e => e.classList.remove('xshot-on')); }''')
            page.evaluate("document.querySelectorAll('[data-xid]').forEach(e => e.removeAttribute('data-xid'));"
                          "document.querySelectorAll('.xpseudo').forEach(e => e.remove())")
            return {'items': items, 'pics': pics, 'info': info}

        for m in manifest:
            kind = 'title' if m['k'] in ('title', 'thankyou', 'section') else 'content'
            slides.append(grab(m['i'], kind))
            print(f'  measured slide {m["i"] + 1:3d}/{len(manifest)}', end='\r', flush=True)
        browser.close()
    print()
    return manifest, slides


# ── the PowerPoint side ──────────────────────────────────────────────────

class Writer:
    """Slides on the layouts of the university master: its logo, rules, footer, page tab, seal and theme (Segoe UI,
    university palette) come from the master; the deck fills the layout placeholders and adds its content on top."""
    official = True          # placeholders keep the layout's geometry and text styles

    def __init__(self, master):
        self.font = None        # text takes the master's theme font; only code is set (MONO)
        self.prs = open_template(master)
        lst = self.prs.slides._sldIdLst                    # drop the master's sample slides
        for sid in list(lst):
            self.prs.part.drop_rel(sid.rId)
            lst.remove(sid)
        self.layouts = {l.name: l for l in self.prs.slide_layouts}
        self.anim = []          # (slide, [(step, shape, paragraph index or None)])
        self.jumps = []         # (shape, target slide index)
        self.warnings = []

    def finish(self, slide, m, data, phs):
        self.emit(slide.shapes, data['items'], data['pics'], phs, slide_index=m['i'])
        for ph in list(slide.placeholders):    # unused placeholders would show "Click to add title" while editing
            t = str(ph.placeholder_format.type)
            base = ph._base_placeholder
            decor = base is not None and b'blipFill' in etree.tostring(base._element)   # the logo or seal as a picture fill
            if decor:
                continue
            if 'PICTURE' in t or (ph.has_text_frame and not ph.text_frame.text.strip()
                                  and not any(k in t for k in ('SLIDE_NUMBER', 'FOOTER', 'DATE'))):
                ph._element.getparent().remove(ph._element)
        if m.get('n'):
            slide.notes_slide.notes_text_frame.text = strip_html(m['n'])
        return slide

    def emit(self, shapes, items, pics, phs, slide_index):
        steps = []
        for it in items:
            t = it['type']
            shape = None
            if t == 'box':
                shape = self.box(shapes, it)
            elif t == 'line':
                x1, y1, x2, y2 = it['pts']
                shape = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
                shape.line.color.rgb = rgb(it['c'])
                shape.line.width = E(it['w'])
            elif t == 'img':
                blob = pics.get(it['xid']) if it.get('xid') else base64.b64decode(it['src'].split(',', 1)[1])
                if blob is None:
                    self.warnings.append(f'slide {slide_index}: picture without data')
                    continue
                shape = shapes.add_picture(io.BytesIO(blob), E(it['x']), E(it['y']), E(it['w']), E(it['h']))
            elif t == 'table':
                shape = self.table(shapes, it)
            elif t == 'text':
                free = [ph for ph in (phs or {}).get(it.get('role') or '', []) if not ph.text_frame.text.strip()
                        and (not self.official or (abs(ph.left / PX - it['x']) < 130 and abs(ph.top / PX - it['y']) < 260))]
                if free:
                    shape = free[0]
                    if not self.official:
                        shape.left, shape.top, shape.width, shape.height = E(it['x']), E(it['y']), E(it['w']), E(it['h'] + 4)
                    self.text(shape, it, inherit=self.official)
                else:
                    shape = shapes.add_textbox(E(it['x']), E(it['y']), E(it['w']), E(it['h']))
                    self.text(shape, it)
            if shape is None:
                continue
            if it.get('jump') is not None and slide_index is not None:
                self.jumps.append((shape, it['jump']))
            if slide_index is not None:
                if it.get('step'):
                    steps.append((it['step'], shape, None))
                elif t == 'text':
                    for k, p in enumerate(it['paras']):
                        if p.get('step'):
                            steps.append((p['step'], shape, k))
        if slide_index is not None and steps:
            self.anim.append((shapes, steps))

    def box(self, shapes, it):
        kind = {'chevron': MSO_SHAPE.CHEVRON, 'pentagon': MSO_SHAPE.PENTAGON, 'diamond': MSO_SHAPE.DIAMOND}.get(it['shape'])
        if kind is None:
            kind = MSO_SHAPE.ROUNDED_RECTANGLE if it.get('radius') else MSO_SHAPE.RECTANGLE
        s = shapes.add_shape(kind, E(it['x']), E(it['y']), E(it['w']), E(it['h']))
        if kind in (MSO_SHAPE.CHEVRON, MSO_SHAPE.PENTAGON) and it.get('notch'):
            s.adjustments[0] = min(1.0, it['notch'] / min(it['w'], it['h']))
        elif kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            s.adjustments[0] = min(0.5, it['radius'] / min(it['w'], it['h']))
        if it.get('fill'):
            s.fill.solid()
            s.fill.fore_color.rgb = rgb(it['fill'])
            self._alpha(s.fill._xPr.find(qn('a:solidFill')), it['fill'])
        else:
            s.fill.background()
        if it.get('line'):
            s.line.color.rgb = rgb(it['line']['c'])
            s.line.width = E(it['line']['w'])
        else:
            s.line.fill.background()
        s.shadow.inherit = False
        s.text_frame.text = ''
        return s

    @staticmethod
    def _alpha(solid, c):
        if solid is not None and len(c) > 3 and c[3] < 0.999:
            clr = solid[0]
            etree.SubElement(clr, qn('a:alpha'), val=str(int(c[3] * 100000)))

    def text(self, shape, it, inherit=False):
        tf = shape.text_frame
        if inherit:                    # a placeholder of the official master: its own geometry, styles and bullets
            for k, p in enumerate(it['paras']):
                para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                self.runs(para, p, inherit=True)
            return
        single = it.get('single') and len(it['paras']) == 1
        tf.word_wrap = not single
        tf.auto_size = MSO_AUTO_SIZE.NONE
        l, t, r, b = it.get('pad') or (0, 0, 0, 0)
        tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = E(l), E(t), E(r), E(b)
        tf.vertical_anchor = MSO_ANCHOR.TOP
        if it.get('cols'):                   # CSS columns -> text box columns
            body_pr = tf._txBody.find(qn('a:bodyPr'))
            body_pr.set('numCol', str(it['cols'][0]))
            body_pr.set('spcCol', str(int(E(it['cols'][1]))))
        if it.get('fill'):
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(it['fill'])
        if it.get('line'):
            shape.line.color.rgb = rgb(it['line']['c'])
            shape.line.width = E(it['line']['w'])
        for k, p in enumerate(it['paras']):
            para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            if k == 0:
                for r_ in list(para.runs):
                    r_._r.getparent().remove(r_._r)
            self.paragraph(para, p, first=(k == 0))

    def paragraph(self, para, p, first):
        para.alignment = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT, 'end': PP_ALIGN.RIGHT,
                          'justify': PP_ALIGN.JUSTIFY}.get(p['align'], PP_ALIGN.LEFT)
        para.line_spacing = Pt(p['lh'] * 0.75)
        if not first and p.get('before', 0) > 0.5:
            para.space_before = Pt(p['before'] * 0.75)
        ppr = para._p.get_or_add_pPr()
        bu = p.get('bullet')
        if bu:
            ppr.set('marL', str(int(E(bu['ind']))))
            ppr.set('indent', str(int(E(-bu['ind']))))
            if bu['kind'] == 'text':             # a CSS marker (A., 1), a dash): its own run before a tab at the indent
                etree.SubElement(ppr, qn('a:buNone'))
                p = dict(p, runs=[dict(bu['run'], t=bu['t'], pre=True), dict(bu['run'], t='\t', pre=True)] + list(p['runs']))
            elif bu['kind'] == 'sq':
                clr = etree.SubElement(ppr, qn('a:buClr'))
                etree.SubElement(clr, qn('a:srgbClr'), val=hexc(bu['c'] or (21, 63, 135)))
                pct = max(25, min(100, round(100 * bu['size'] / p['sz'] / 0.62)))
                etree.SubElement(ppr, qn('a:buSzPct'), val=str(pct * 1000))
                etree.SubElement(ppr, qn('a:buFont'), typeface='Arial')
                etree.SubElement(ppr, qn('a:buChar'), char='■')
            else:
                etree.SubElement(ppr, qn('a:buFont'), typeface='+mj-lt')
                etree.SubElement(ppr, qn('a:buAutoNum'), type='arabicPeriod', startAt=str(bu['start']))
        else:
            ppr.set('marL', '0')
            ppr.set('indent', '0')
            etree.SubElement(ppr, qn('a:buNone'))
        self.runs(para, p)

    def runs(self, para, p, inherit=False):
        """inherit: only what differs from the paragraph (a kicker's colour, a bold word); size and font come from the layout."""
        runs = self.clean_runs(p['runs'])
        if not runs:
            if not inherit:
                para.font.size = Pt(p['sz'] * 0.75)
            return
        for r in runs:
            if r.get('br'):
                para.add_line_break()
                continue
            run = para.add_run()
            run.text = r['t'].upper() if r.get('up') else r['t']
            f = run.font
            if inherit:
                if r['b'] != p.get('b'):
                    f.bold = r['b']
                if r['i']:
                    f.italic = True
                if r.get('mono'):
                    f.name = MONO
                if r.get('c') and (r.get('op', 1) < 0.98 or (p.get('c') and hexc(r['c']) != hexc(p['c']))):
                    f.color.rgb = rgb(r['c'])
            else:
                f.size = Pt(r['sz'] * 0.75)
                f.bold = r['b']
                f.italic = r['i']
                if r.get('mono'):
                    f.name = MONO
                elif self.font:
                    f.name = self.font
                if r.get('c'):
                    f.color.rgb = rgb(r['c'])
            rpr = run._r.get_or_add_rPr()
            alpha = (r['c'][3] if r.get('c') and len(r['c']) > 3 else 1) * r.get('op', 1)
            if alpha < 0.98:
                if not r.get('c'):
                    f.color.rgb = RGBColor(0, 0, 0)
                clr = rpr.find(qn('a:solidFill'))
                if clr is not None and len(clr):
                    etree.SubElement(clr[0], qn('a:alpha'), val=str(int(alpha * 100000)))
            if r.get('ls'):
                rpr.set('spc', str(int(round(r['ls'] * 0.75 * 100))))
            if r.get('va'):
                rpr.set('baseline', '30000' if r['va'] == 'sup' else '-25000')

    @staticmethod
    def clean_runs(runs):
        """Collapse HTML whitespace like the browser does; keep it in <pre> text; split newlines into breaks."""
        out = []
        for r in runs:
            if r.get('br'):
                out.append(r)
                continue
            if r.get('pre'):
                lines = r['t'].split('\n')
                for k, line in enumerate(lines):
                    if k:
                        out.append({'br': True})
                    if line:
                        out.append(dict(r, t=line))
                continue
            t = re.sub(r'\s+', ' ', r['t'])
            if t:
                out.append(dict(r, t=t))
        # trim at the start, the end, and around breaks; collapse double spaces across runs
        prev_space = True
        for r in out:
            if r.get('br'):
                prev_space = True
                continue
            if r.get('pre'):
                prev_space = r['t'].endswith(('\t', ' '))
                continue
            if prev_space:
                r['t'] = r['t'].lstrip(' ')
            prev_space = r['t'].endswith(' ') or (not r['t'] and prev_space)
        for k in range(len(out) - 1, -1, -1):
            if out[k].get('br') or out[k].get('pre'):
                if out[k].get('br') and k == len(out) - 1:
                    out.pop()
                    continue
                break
            out[k]['t'] = out[k]['t'].rstrip(' ')
            if out[k]['t']:
                break
        return [r for r in out if r.get('br') or r['t']]

    def table(self, shapes, it):
        rows = it['rows']
        ncols = max(sum(c['span'] for c in row) for row in rows)
        gf = shapes.add_table(len(rows), ncols, E(it['x']), E(it['y']), E(it['w']), E(it['h']))
        tbl = gf.table
        tblpr = tbl._tbl.tblPr
        for k in ('firstRow', 'bandRow'):
            tblpr.set(k, '0')
        style = tblpr.find(qn('a:tableStyleId'))
        if style is not None:
            style.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'      # "No Style, No Grid"
        first = rows[0]
        for k, c in enumerate(first[:ncols]):
            tbl.columns[k].width = E(c['r']['w'])
        for k, row in enumerate(rows):
            tbl.rows[k].height = E(row[0]['r']['h'])
            for j, c in enumerate(row[:ncols]):
                cell = tbl.cell(k, j)
                l, t, r, b = c['pad']
                cell.margin_left, cell.margin_top, cell.margin_right, cell.margin_bottom = E(l), E(t), E(r), E(b)
                cell.vertical_anchor = MSO_ANCHOR.TOP
                if c.get('fill'):
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(c['fill'])
                else:
                    cell.fill.background()
                tcpr = cell._tc.get_or_add_tcPr()
                for side, tagname in (('top', 'a:lnT'), ('bottom', 'a:lnB')):
                    if c.get(side):
                        ln = etree.Element(qn(tagname), w=str(int(E(c[side]['w']))))
                        sf = etree.SubElement(ln, qn('a:solidFill'))
                        etree.SubElement(sf, qn('a:srgbClr'), val=hexc(c[side]['c']))
                        tcpr.insert(0 if side == 'top' else len(tcpr.findall(qn('a:lnT'))), ln)
                para = cell.text_frame.paragraphs[0]
                self.paragraph(para, c['p'], first=True)
        return gf

    # step builds -> appear on click, in step order; paragraphs of one list animate on their own
    def animations(self):
        for shapes, steps in self.anim:
            slide_el = shapes._spTree.getparent().getparent()
            by_step = {}
            for step, shape, k in steps:
                by_step.setdefault(step, []).append((shape, k))
            clicks = []
            ctn = [2]

            def nid():
                ctn[0] += 1
                return ctn[0]
            for n, step in enumerate(sorted(by_step)):
                effects = []
                for j, (shape, k) in enumerate(by_step[step]):
                    target = (f'<p:tgtEl><p:spTgt spid="{shape.shape_id}"/></p:tgtEl>' if k is None else
                              f'<p:tgtEl><p:spTgt spid="{shape.shape_id}"><p:txEl><p:pRg st="{k}" end="{k}"/></p:txEl></p:spTgt></p:tgtEl>')
                    node = 'clickEffect' if j == 0 else 'withEffect'
                    effects.append(f'''<p:par><p:cTn id="{nid()}" presetID="1" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">
<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr><p:cTn id="{nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>
{target}<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set></p:childTnLst></p:cTn></p:par>''')
                clicks.append(f'''<p:par><p:cTn id="{nid()}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="{nid()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{''.join(effects)}</p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn></p:par>''')
            seen, blds = set(), []
            for step, shape, k in steps:
                if shape.shape_id in seen or shape._element.tag != qn('p:sp') or not shape.has_text_frame:
                    continue                    # bldP is only for shapes with text; pictures and lines animate without one
                seen.add(shape.shape_id)
                build = ' build="p"' if any(sh.shape_id == shape.shape_id and kk is not None for _, sh, kk in steps) else ''
                blds.append(f'<p:bldP spid="{shape.shape_id}" grpId="0"{build}/>')
            # an empty bldLst is invalid (PowerPoint repairs the file): a slide that only builds pictures has none
            bld_lst = f'<p:bldLst>{"".join(blds)}</p:bldLst>' if blds else ''
            xml = f'''<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{''.join(clicks)}</p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
</p:childTnLst></p:cTn></p:par></p:tnLst>{bld_lst}</p:timing>'''
            seq = iter(range(1, 10 ** 6))      # PowerPoint wants the cTn ids in document order
            xml = re.sub(r'<p:cTn id="\d+"', lambda mo: f'<p:cTn id="{next(seq)}"', xml)
            timing = etree.fromstring(xml)
            ext = slide_el.find(qn('p:extLst'))
            if ext is not None:
                ext.addprevious(timing)
            else:
                slide_el.append(timing)

    def links(self):
        slides = list(self.prs.slides)
        for shape, target in self.jumps:
            if 0 <= target < len(slides):
                shape.click_action.target_slide = slides[target]

    def sections(self, manifest):
        """PowerPoint sections from the deck's groups (Start, the meta() sections, Backup), replacing any of the template."""
        P14 = 'http://schemas.microsoft.com/office/powerpoint/2010/main'
        URI = '{521415D9-36F7-43E2-AB2F-B90AF26B5E84}'
        pres = self.prs.part._element
        ext_lst = pres.find(qn('p:extLst'))
        if ext_lst is None:
            ext_lst = etree.SubElement(pres, qn('p:extLst'))
        for ext in ext_lst.findall(qn('p:ext')):
            if ext.get('uri') == URI:
                ext_lst.remove(ext)
        ids = [sid.get('id') for sid in self.prs.slides._sldIdLst]
        groups = []
        for m, sid in zip(manifest, ids):
            name = m.get('g') or 'Start'
            if not groups or groups[-1][0] != name:
                groups.append((name, []))
            groups[-1][1].append(sid)
        ext = etree.SubElement(ext_lst, qn('p:ext'), uri=URI)
        lst = etree.SubElement(ext, f'{{{P14}}}sectionLst', nsmap={'p14': P14})
        for k, (name, sids) in enumerate(groups):
            sec = etree.SubElement(lst, f'{{{P14}}}section', name=name, id='{%08X-0000-4000-8000-%012X}' % (k + 1, k + 1))
            sl = etree.SubElement(sec, f'{{{P14}}}sldIdLst')
            for sid in sids:
                etree.SubElement(sl, f'{{{P14}}}sldId', id=sid)
        ext_lst.insert(0, ext)   # PowerPoint lists the section extension first

    def save(self, out, manifest=()):
        self.animations()
        self.links()
        if manifest:
            self.sections(manifest)
        self.prs.save(out)


UNI_MASTER = 'PPT-Master-Universitaet-Wuerzburg.potx'
SLIDENUM = ('<a:fld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}" '
            'type="slidenum"><a:t>&#8249;#&#8250;</a:t></a:fld>')


def find_master(path=None, deck=None):
    """The university .potx. It is not part of the skill (the repository is public): --master, else $D3_UNI_MASTER,
    else the first PPT-Master-Universitaet-Wuerzburg.potx in the deck's folder or one of its parents, else d3deck/office/."""
    cands = [path, os.environ.get('D3_UNI_MASTER')]
    d = os.path.dirname(os.path.abspath(deck)) if deck else os.getcwd()
    while True:
        cands.append(os.path.join(d, UNI_MASTER))
        up = os.path.dirname(d)
        if up == d:
            break
        d = up
    cands.append(os.path.join(HERE, 'office', UNI_MASTER))
    for c in cands:
        if c and os.path.exists(c):
            return c
    sys.exit(f'{UNI_MASTER} not found. The uni look exports onto the university PowerPoint master: put the .potx into the '
             f'talk folder or a parent folder, set D3_UNI_MASTER, or pass --master <path>')


def open_template(path):
    """python-pptx opens .pptx only: rewrite the template's main content type in memory."""
    import zipfile
    src, buf = zipfile.ZipFile(path), io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == '[Content_Types].xml':
                data = data.replace(b'presentationml.template.main+xml', b'presentationml.presentation.main+xml')
            elif item.filename == 'ppt/viewProps.xml':
                # the template was saved in master view: open the deck on its slides instead
                data = re.sub(rb'\slastView="[^"]*"', b'', data)
            out.writestr(item, data)
    buf.seek(0)
    return Presentation(buf)


class MasterWriter(Writer):
    """Picks the university layout for each slide and maps the deck's frame onto its placeholders."""

    def pick(self, m, data):
        info, items = data['info'], data['items']
        if m['k'] == 'title':
            return UNI_COVER.get(info.get('cover') or 'navy', 'Deckblatt 1')
        if m['k'] == 'section':
            return 'Trennseite 2'
        if m['k'] == 'thankyou':
            return 'Abschlussseite'
        if info.get('agenda'):
            return 'Agenda 1'
        if info.get('refs') and info.get('title') == 'Quellenangaben':
            return 'Quellenangaben'
        if info.get('take'):
            return 'Inhaltsfolie 2'
        # the body only: the header (D3 logo, section tracker) and the title do not decide the layout
        rest = [it for it in items if it.get('role') not in ('title', 'cobrand') and it.get('y', 999) > 190]
        if (len(rest) == 1 and rest[0]['type'] == 'text' and not info.get('refs') and rest[0]['paras'][0]['sz'] >= 20
                and all(p.get('bullet') and p['bullet']['kind'] == 'sq' for p in rest[0]['paras'])):
            rest[0]['role'] = 'body'                        # a plain list goes into the content placeholder
            return 'Inhaltsfolie 1'
        return 'Blanko 1'

    def add_slide(self, m, data):
        name = self.pick(m, data)
        layout = self.layouts[name]
        slide = self.prs.slides.add_slide(layout)
        by_idx = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
        table = UNI_PH['Deckblatt' if name.startswith('Deckblatt') else name] if (name.startswith('Deckblatt') or name in UNI_PH) else {}
        phs = {role: [by_idx[i] for i in idxs if i in by_idx] for role, idxs in table.items()}
        if m['k'] == 'thankyou':
            self.contact(data['items'], phs)
        info = data['info']
        if info.get('num') or info.get('foot'):
            for lph in layout.placeholders:
                t = str(lph.placeholder_format.type)
                if 'SLIDE_NUMBER' in t and info.get('num'):
                    slide.shapes.clone_placeholder(lph)
                    ph = slide.placeholders[lph.placeholder_format.idx]
                    p = ph.text_frame.paragraphs[0]
                    p._p.append(etree.fromstring(SLIDENUM))
                elif 'FOOTER' in t and info.get('foot'):
                    slide.shapes.clone_placeholder(lph)
                    ph = slide.placeholders[lph.placeholder_format.idx]
                    ph.text_frame.text = info['foot']
                    if info.get('footw'):
                        ph.left, ph.top, ph.width, ph.height = lph.left, lph.top, E(info['footw']), lph.height
        slide = self.finish(slide, m, data, phs)
        if name.startswith('Deckblatt'):
            for ph in slide.placeholders:
                if ph.placeholder_format.idx == 0:
                    ph.text_frame._txBody.find(qn('a:bodyPr')).set('anchor', 'b')
        return slide

    @staticmethod
    def contact(items, phs):
        """Abschlussseite: name, phone (after the printed 'Tel.') and mail each have their own placeholder."""
        rows = [it for it in items if it.get('role') == 'contact']
        slots = phs.get('contact', [])
        if len(slots) < 3:
            return
        name, tel, mail = slots[:3]
        phs['contact'] = []
        phs['c_name'], phs['c_tel'], phs['c_mail'] = [name], [tel], [mail]
        for it in rows:
            text = ''.join(r.get('t', '') for p in it['paras'] for r in p['runs']).strip()
            if '@' in text:
                it['role'] = 'c_mail'
            elif re.match(r'(?i)tel\.?\s', text) or re.fullmatch(r'[+\d][\d\s/()-]{5,}', text):
                it['role'] = 'c_tel'
                for p in it['paras']:
                    for r in p['runs']:
                        if r.get('t'):
                            r['t'] = re.sub(r'(?i)^\s*tel\.?\s*', '', r['t'])
            else:
                it['role'] = 'c_name'


def export(deck, out=None, notes=True, master=None):
    master = find_master(master, deck)          # before the browser starts: a missing master fails fast
    manifest, slides = measure(deck)
    w = MasterWriter(master)
    for m, data in zip(manifest, slides):
        if not notes:
            m = dict(m, n='')
        w.add_slide(m, data)
    out = out or os.path.splitext(deck)[0] + '.pptx'
    w.save(out, manifest)
    for msg in w.warnings:
        print('warning:', msg)
    pics = sum(1 for d in slides for it in d['items'] if it['type'] == 'img')
    print(f'{len(manifest)} slides · {pics} pictures · {len(w.anim)} slides with builds → {out}')
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('deck', help='the built deck, dist/<name>.html')
    ap.add_argument('--out', default=None, help='output path (default: next to the deck, .pptx)')
    ap.add_argument('--master', default=None, help='path to PPT-Master-Universitaet-Wuerzburg.potx '
                    '(default: $D3_UNI_MASTER, then the talk folder and its parents)')
    ap.add_argument('--no-notes', action='store_true', help='leave the speaker notes out')
    a = ap.parse_args(argv)
    export(a.deck, a.out, notes=not a.no_notes, master=a.master)


if __name__ == '__main__':
    main()
